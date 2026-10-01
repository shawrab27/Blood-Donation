import re
import csv
import os
import unicodedata
from collections import Counter
import openpyxl
from api.models import District, Institution, InstitutionAlias

DISTRICT_ALIAS_MAP = {
    'netrakona': 'netrokona',
    'nawabganj': 'chapainawabganj',
    'brahamanbaria': 'brahmanbaria',
    'coxs bazar': "cox's bazar",
    'jhalokati': 'jhalokathi',
    'maulvibazar': 'moulvibazar',
    'norail': 'narail',
}


def normalize_val(val):
    if val is None:
        return None
    s = unicodedata.normalize('NFC', str(val)).strip()
    s = re.sub(r'\s+', ' ', s)
    return s if s else None


def get_district_cache():
    return {d.name.strip().lower(): d for d in District.objects.exclude(name__contains='\u2502')}


def parse_banbeis_file(file_path, file_name, itype, district_map):
    header_row = 2 if file_name in ['college.xlsx', 'school_list.xlsx'] else 1
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    col_map = {}
    for idx, h in enumerate(headers, start=1):
        if not h:
            continue
        hl = str(h).strip().lower()
        if hl in ['institute_name', 'institution_name', 'name']:
            col_map['name'] = idx
        elif hl in ['eiin_no', 'eiin']:
            col_map['eiin'] = idx
        elif hl in ['division_name', 'division']:
            col_map['division'] = idx
        elif hl in ['district_name', 'district']:
            col_map['district'] = idx
        elif hl in ['thana_name', 'thana', 'upazila_name', 'upazila']:
            col_map['thana'] = idx

    if 'name' not in col_map:
        raise ValueError(f"Ambiguous or missing name column in {file_name}")

    rows_read = 0
    records = []
    rejected = []
    unmatched_districts = Counter()

    for r in range(header_row + 1, ws.max_row + 1):
        raw_name = ws.cell(r, col_map['name']).value if 'name' in col_map else None
        raw_eiin = ws.cell(r, col_map['eiin']).value if 'eiin' in col_map else None
        raw_div = ws.cell(r, col_map['division']).value if 'division' in col_map else None
        raw_dist = ws.cell(r, col_map['district']).value if 'district' in col_map else None
        raw_thana = ws.cell(r, col_map['thana']).value if 'thana' in col_map else None

        if all(v is None or str(v).strip() == '' for v in [raw_name, raw_eiin, raw_div, raw_dist, raw_thana]):
            continue

        rows_read += 1
        name = normalize_val(raw_name)
        if not name:
            rejected.append((r, "Empty institution name"))
            continue

        eiin = normalize_val(raw_eiin)
        div_name = normalize_val(raw_div)
        dist_name = normalize_val(raw_dist)
        upazila_name = normalize_val(raw_thana)

        clean_dist = dist_name.lower() if dist_name else ''
        matched_district = district_map.get(clean_dist)
        if not matched_district and clean_dist in DISTRICT_ALIAS_MAP:
            matched_district = district_map.get(DISTRICT_ALIAS_MAP[clean_dist])

        if dist_name and not matched_district:
            unmatched_districts[dist_name] += 1

        dedupe_key = ('eiin', eiin) if eiin else ('name_dist', name.lower(), (dist_name or '').lower())
        records.append({
            'dedupe_key': dedupe_key,
            'name': name,
            'eiin': eiin,
            'institution_type': itype,
            'district': matched_district,
            'division_name': div_name,
            'district_name': dist_name,
            'upazila_name': upazila_name,
            'source_file': file_name,
        })

    return {
        'rows_read': rows_read,
        'records': records,
        'rejected': rejected,
        'unmatched_districts': unmatched_districts,
    }


def load_institution_aliases_from_csv(csv_path, stdout=None):
    if not os.path.exists(csv_path):
        if stdout:
            stdout.write(f"Aliases file not found at {csv_path}, skipping.")
        return 0, 0

    institutions_by_eiin = {inst.eiin: inst for inst in Institution.objects.all() if inst.eiin}
    loaded = 0
    skipped = 0

    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            alias = normalize_val(row.get('alias'))
            eiin = normalize_val(row.get('eiin'))
            note = normalize_val(row.get('source_note')) or ''

            if not alias or not eiin:
                skipped += 1
                continue

            inst = institutions_by_eiin.get(eiin)
            InstitutionAlias.objects.update_or_create(
                alias=alias,
                eiin=eiin,
                defaults={
                    'institution': inst,
                    'source_note': note,
                }
            )
            loaded += 1

    if stdout:
        stdout.write(f"Loaded {loaded} institution aliases (skipped {skipped} invalid rows).")
    return loaded, skipped
